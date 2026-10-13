"""Optional editable captions added after the upstream image-only PPTX export."""
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
import argparse
p=argparse.ArgumentParser();p.add_argument('input',type=Path);p.add_argument('--output',type=Path,required=True);a=p.parse_args()
prs=Presentation(a.input)
titles=[('FIELD NOTES','A fictional small-batch coffee story'),('ROAST & CRAFT','Textures, materials and a shared visual language'),('A PLACE TO PAUSE','A reusable community presentation')]
for slide,(title,subtitle) in zip(prs.slides,titles):
    panel=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(.35), Inches(.8), Inches(5.1), Inches(2.1));panel.fill.solid();panel.fill.fore_color.rgb=RGBColor(250,247,240);panel.line.fill.background()
    box=slide.shapes.add_textbox(Inches(.65),Inches(1), Inches(4.4), Inches(1.55));tf=box.text_frame;tf.clear()
    para=tf.paragraphs[0];para.text=title;para.font.size=Pt(32);para.font.bold=True;para.font.color.rgb=RGBColor(29,29,29)
    para=tf.add_paragraph();para.text=subtitle;para.font.size=Pt(17);para.font.color.rgb=RGBColor(90,83,73)
prs.save(a.output)
check=Presentation(a.output);assert len(check.slides)==3
assert sum(1 for s in check.slides for sh in s.shapes if sh.has_text_frame and sh.text)==3
print('Saved three-slide copy with native editable caption boxes; original images unchanged')
